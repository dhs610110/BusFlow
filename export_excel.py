import sqlite3
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from config import ROUTES, STATIONS_5001A, STATIONS_5001B


DB_PATH = "data/busflow.db"
OUTPUT_PATH = "data/busflow.xlsx"


# 노선 ID → 사람이 읽을 수 있는 노선명
ROUTE_NAME_MAP = {
    ROUTES["5001A"]: "5001A",
    ROUTES["5001B"]: "5001B",
}


# 정류장 ID → 정류장 정보
STATION_MAP = {}

for station in STATIONS_5001A:
    STATION_MAP[station["id"]] = {
        "seq": station["seq"],
        "name": station["name"],
        "route": "5001A",
    }

for station in STATIONS_5001B:
    STATION_MAP[station["id"]] = {
        "seq": station["seq"],
        "name": station["name"],
        "route": "5001B",
    }


def get_data():
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()

    cursor.execute("""
        SELECT
            opr_ymd,
            dow_nm,
            route_id,
            station_id,
            station_seq,
            time_zone,
            congestion
        FROM congestion
        ORDER BY
            opr_ymd,
            route_id,
            station_seq,
            time_zone
    """)

    rows = cursor.fetchall()
    conn.close()

    return rows


def style_sheet(ws):
    # 첫 줄 고정
    ws.freeze_panes = "A2"

    # 헤더 디자인
    header_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7"
    )

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    # 열 너비 자동 조절
    for column_cells in ws.columns:
        max_length = 0
        column_letter = get_column_letter(column_cells[0].column)

        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, len(value))

        ws.column_dimensions[column_letter].width = min(
            max_length + 3,
            35
        )

    ws.auto_filter.ref = ws.dimensions


def create_raw_sheet(wb, rows):
    ws = wb.active
    ws.title = "원본데이터"

    headers = [
        "날짜",
        "요일",
        "노선",
        "정류장순서",
        "정류장명",
        "정류장ID",
        "시간대",
        "혼잡도",
    ]

    ws.append(headers)

    for row in rows:
        (
            opr_ymd,
            dow_nm,
            route_id,
            station_id,
            station_seq,
            time_zone,
            congestion,
        ) = row

        route_name = ROUTE_NAME_MAP.get(route_id, route_id)

        station_info = STATION_MAP.get(station_id)

        if station_info:
            station_name = station_info["name"]
            station_seq = station_info["seq"]
        else:
            station_name = "미등록 정류장"

        formatted_date = (
            f"{opr_ymd[:4]}-{opr_ymd[4:6]}-{opr_ymd[6:8]}"
        )

        ws.append([
            formatted_date,
            dow_nm,
            route_name,
            station_seq,
            station_name,
            station_id,
            time_zone,
            congestion,
        ])

    style_sheet(ws)


def create_average_sheet(wb, rows, route_name, stations):
    ws = wb.create_sheet(
        f"{route_name}_요일시간대평균"
    )

    headers = [
        "요일",
        "시간대",
        "정류장순서",
        "정류장명",
        "평균혼잡도",
        "데이터수",
    ]

    ws.append(headers)

    target_route_id = ROUTES[route_name]

    target_station_ids = {
        station["id"] for station in stations
    }

    station_lookup = {
        station["id"]: station
        for station in stations
    }

    grouped = defaultdict(list)

    for row in rows:
        (
            opr_ymd,
            dow_nm,
            route_id,
            station_id,
            station_seq,
            time_zone,
            congestion,
        ) = row

        if route_id != target_route_id:
            continue

        # 우리가 분석하기로 한 정류장만
        if station_id not in target_station_ids:
            continue

        key = (
            dow_nm,
            time_zone,
            station_id,
        )

        grouped[key].append(congestion)

    # API에서 실제 어떤 형태로 요일이 저장됐든 정렬 가능하도록
    day_order = {
        "월": 1,
        "월요일": 1,
        "화": 2,
        "화요일": 2,
        "수": 3,
        "수요일": 3,
        "목": 4,
        "목요일": 4,
        "금": 5,
        "금요일": 5,
        "토": 6,
        "토요일": 6,
        "일": 7,
        "일요일": 7,
    }

    sorted_keys = sorted(
        grouped.keys(),
        key=lambda key: (
            day_order.get(key[0], 99),
            str(key[1]),
            station_lookup[key[2]]["seq"],
        )
    )

    for dow_nm, time_zone, station_id in sorted_keys:

        congestion_values = grouped[
            (
                dow_nm,
                time_zone,
                station_id,
            )
        ]

        average = sum(congestion_values) / len(
            congestion_values
        )

        station = station_lookup[station_id]

        ws.append([
            dow_nm,
            time_zone,
            station["seq"],
            station["name"],
            round(average, 1),
            len(congestion_values),
        ])

    style_sheet(ws)


def main():
    print("DB 데이터 읽는 중...")

    rows = get_data()

    print("현재 저장된 데이터:", len(rows), "개")

    wb = Workbook()

    create_raw_sheet(wb, rows)

    create_average_sheet(
        wb,
        rows,
        "5001A",
        STATIONS_5001A
    )

    create_average_sheet(
        wb,
        rows,
        "5001B",
        STATIONS_5001B
    )

    wb.save(OUTPUT_PATH)

    print()
    print("============================")
    print("Excel 생성 완료")
    print(OUTPUT_PATH)
    print("============================")


if __name__ == "__main__":
    main()